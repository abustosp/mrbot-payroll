"""Export tabular de todos los modulos en csv, xlsx y json."""

from __future__ import annotations

import io
from collections.abc import Generator
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from app.database import Base
from app.database.session import get_db
from app.main import app
from scripts.seed_demo import seed

MODULES = [
    "/api/v1/employees/export",
    "/api/v1/companies/export",
    "/api/v1/agreements/export",
    "/api/v1/concepts/export",
    "/api/v1/formulas/export",
    "/api/v1/payroll-runs/export",
]


def _client(tmp_path: Path) -> tuple[TestClient, dict]:
    engine = create_engine(f"sqlite:///{tmp_path / 'exports.sqlite'}")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)
    with factory.begin() as session:
        demo = seed(session)
        context = {"company_id": demo["company"].id}

    def override_get_db() -> Generator[Session, None, None]:
        with factory() as session:
            yield session

    app.dependency_overrides[get_db] = override_get_db
    return TestClient(app), context


def test_todos_los_modulos_exportan_csv_xlsx_json(tmp_path: Path) -> None:
    client, _context = _client(tmp_path)
    try:
        login = client.post(
            "/api/v1/auth/login",
            json={"email": "admin@example.com", "password": "admin"},
        )
        assert login.status_code == 200

        for path in MODULES:
            csv_response = client.get(path, params={"format": "csv"})
            assert csv_response.status_code == 200, (path, csv_response.text)
            assert "text/csv" in csv_response.headers["content-type"]
            assert 'attachment; filename="' in csv_response.headers["content-disposition"]

            xlsx_response = client.get(path, params={"format": "xlsx"})
            assert xlsx_response.status_code == 200, (path, xlsx_response.text)
            sheet = load_workbook(io.BytesIO(xlsx_response.content)).active
            assert sheet is not None and sheet.max_row >= 1

            json_response = client.get(path, params={"format": "json"})
            assert json_response.status_code == 200, (path, json_response.text)
            assert isinstance(json_response.json(), list)

        bad = client.get(MODULES[0], params={"format": "pdf"})
        assert bad.status_code == 400
    finally:
        app.dependency_overrides.pop(get_db, None)


def test_resultados_de_liquidacion_exportan_en_tres_formatos(tmp_path: Path) -> None:
    client, context = _client(tmp_path)
    try:
        login = client.post(
            "/api/v1/auth/login",
            json={"email": "admin@example.com", "password": "admin"},
        )
        assert login.status_code == 200

        created = client.post(
            "/api/v1/payroll-runs",
            json={
                "company_id": context["company_id"],
                "code": "export-demo",
                "period_start": "2026-06-01",
                "period_end": "2026-06-30",
                "run_type": "MONTHLY",
            },
        )
        assert created.status_code == 201, created.text
        run_id = created.json()["id"]
        calculated = client.post(f"/api/v1/payroll-runs/{run_id}/calculate")
        assert calculated.status_code == 200, calculated.text

        path = f"/api/v1/payroll-runs/{run_id}/results/export"
        csv_response = client.get(path, params={"format": "csv"})
        assert csv_response.status_code == 200, csv_response.text
        header = csv_response.text.lstrip("\ufeff").splitlines()[0]
        assert "gross" in header and "net" in header

        json_response = client.get(path, params={"format": "json"})
        assert json_response.status_code == 200, json_response.text
        payload = json_response.json()
        assert len(payload) == 2
        assert all(float(row["net"]) > 0 for row in payload)

        xlsx_response = client.get(path, params={"format": "xlsx"})
        assert xlsx_response.status_code == 200, xlsx_response.text
        sheet = load_workbook(io.BytesIO(xlsx_response.content)).active
        assert sheet is not None and sheet.max_row == 3
    finally:
        app.dependency_overrides.pop(get_db, None)

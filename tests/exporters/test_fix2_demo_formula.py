"""Fix2 exporters: V2-10 (sello DEMO) + V2-12 (prefijo ' en F931)."""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal
from types import SimpleNamespace

from openpyxl import load_workbook
from pypdf import PdfReader

from app.exporters.excel.f931 import build_f931_workbook
from app.exporters.pdf.receipts import DEMO_SEAL, render_receipt_html, render_receipt_pdf


def _context(**overrides) -> dict:
    base: dict = {
        "receipt_number": "R-2026-09-0001",
        "issue_date": date(2026, 9, 30),
        "period": "Septiembre 2026",
        "employer": {"name": "Andes Logística S.A.", "cuit": "30-71234567-8"},
        "employee": {"name": "María Gómez", "cuil": "27-30111222-4"},
        "items": [
            {
                "code": "SUELDO_BASICO",
                "name": "Sueldo básico",
                "kind": "remunerative",
                "amount": "1850000.00",
            }
        ],
        "net": "1850000.00",
    }
    base.update(overrides)
    return base


def test_v2_10_html_demo_lleva_sello_y_normal_no() -> None:
    html_demo = render_receipt_html(_context(is_demo=True))
    assert DEMO_SEAL in html_demo
    assert '<div class="demo-banner"' in html_demo

    html_normal = render_receipt_html(_context())
    assert DEMO_SEAL not in html_normal
    assert '<div class="demo-banner"' not in html_normal


def test_v2_10_alias_demo_mode_tambien_activa_sello() -> None:
    assert DEMO_SEAL in render_receipt_html(_context(demo_mode=True))


def test_v2_10_pdf_demo_lleva_sello_extraible() -> None:
    pdf = render_receipt_pdf(_context(is_demo=True))
    assert pdf.startswith(b"%PDF")
    text = "\n".join(
        page.extract_text() or "" for page in PdfReader(io.BytesIO(pdf)).pages
    )
    assert DEMO_SEAL in text


def test_v2_12_f931_prefija_comilla_y_no_deja_formulas() -> None:
    run = SimpleNamespace(
        id=7,
        code="LIQ-2026-06",
        period_start=date(2026, 6, 1),
        period_end=date(2026, 6, 30),
        status="CALCULATED",
        run_type="MONTHLY",
    )
    company = SimpleNamespace(legal_name="Empresa", cuit="30712345678")
    employee = SimpleNamespace(
        id=1,
        first_name="=1+1",
        last_name="+Perez",
        cuil="20123456786",
        employee_code="@E1",
    )
    result = SimpleNamespace(
        id=10,
        payroll_run_id=7,
        employee_id=1,
        remunerative=Decimal("100.00"),
        non_remunerative=Decimal("0"),
        deductions=Decimal("0"),
        contributions=Decimal("0"),
        gross=Decimal("100.00"),
        net=Decimal("100.00"),
        employer_cost=Decimal("100.00"),
    )
    item = SimpleNamespace(
        id=1,
        payroll_result_id=10,
        employee_id=1,
        code="=BASICO",
        description="-Sueldo",
        nature="REMUNERATIVE",
        quantity=Decimal("1"),
        unit_value=None,
        amount=Decimal("100.00"),
    )
    payload = build_f931_workbook(run, [result], [item], [employee], company)
    workbook = load_workbook(io.BytesIO(payload), data_only=False)
    try:
        texts = [
            str(cell.value)
            for ws in workbook.worksheets
            for row in ws.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value
        ]
        for risky in ("=1+1", "+Perez", "@E1", "=BASICO", "-Sueldo"):
            assert any(
                t.startswith("'") and risky in t for t in texts
            ), risky
        assert all(
            cell.data_type != "f"
            for ws in workbook.worksheets
            for row in ws.iter_rows()
            for cell in row
            if cell.value is not None
        )
    finally:
        workbook.close()

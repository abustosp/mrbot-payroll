"""Descarga tabular de un modulo en csv, xlsx o json."""

from __future__ import annotations

import pytest

from app.exporters.tabular import normalize_format, render_table, table_response

ROWS = [
    {"legajo": "A1", "nombre": "Ana", "neto": "37284.45", "activo": True},
    {"legajo": "A2", "nombre": "Luis", "neto": None, "activo": False},
]
COLUMNS = ["legajo", "nombre", "neto", "activo"]


def test_rechaza_formato_desconocido() -> None:
    with pytest.raises(ValueError, match="Formato no soportado"):
        normalize_format("pdf")


def test_csv_con_bom_y_encabezado() -> None:
    raw = render_table(ROWS, COLUMNS, "csv")
    text = raw.decode("utf-8-sig")
    lines = text.strip().splitlines()
    assert lines[0] == "legajo,nombre,neto,activo"
    assert lines[1] == "A1,Ana,37284.45,True"


def test_xlsx_con_encabezado_y_celdas() -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    sheet = load_workbook(BytesIO(render_table(ROWS, COLUMNS, "xlsx"))).active
    assert sheet is not None
    assert [cell.value for cell in sheet[1]] == COLUMNS
    assert [cell.value for cell in sheet[2]] == ["A1", "Ana", "37284.45", True]


def test_json_es_lista_de_objetos() -> None:
    import json

    # Las celdas vacias viajan como "" en los tres formatos.
    assert json.loads(render_table(ROWS, COLUMNS, "json")) == [
        {"legajo": "A1", "nombre": "Ana", "neto": "37284.45", "activo": True},
        {"legajo": "A2", "nombre": "Luis", "neto": "", "activo": False},
    ]


def test_spreadsheets_keep_untrusted_text_as_text_and_json_unchanged() -> None:
    import csv
    import json
    from decimal import Decimal
    from io import BytesIO, StringIO

    from openpyxl import load_workbook

    values = ["=1+1", "+1+1", "-1+1", "@SUM(A1)", " \t=1+1"]
    rows = [{"name": value, "amount": Decimal("-12.50")} for value in values]
    columns = ["name", "amount"]
    csv_rows = list(csv.reader(StringIO(render_table(rows, columns, "csv").decode("utf-8-sig"))))
    assert [row[0] for row in csv_rows[1:]] == ["'" + value for value in values]
    assert all(row[1] == "-12.50" for row in csv_rows[1:])
    sheet = load_workbook(BytesIO(render_table(rows, columns, "xlsx"))).active
    assert sheet is not None
    assert [row[0].value for row in list(sheet.rows)[1:]] == ["'" + value for value in values]
    assert all(row[0].data_type == "s" and row[1].value == -12.5 for row in list(sheet.rows)[1:])
    assert [row["name"] for row in json.loads(render_table(rows, columns, "json"))] == values


def test_download_filename_handles_unicode_quotes_and_newlines() -> None:
    from urllib.parse import unquote

    name = 'liquidación-"\r\nX-Injected: yes/extra'
    response = table_response([], ["name"], name, "csv")
    header = response.headers["content-disposition"]
    assert "\r" not in header and "\n" not in header
    assert unquote(header.split("filename*=UTF-8''", 1)[1]) == name + "-export.csv"

from __future__ import annotations

import io
from collections.abc import Iterator
from datetime import date
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker

from app.database import Base
from app.database.models import (
    Company,
    Employee,
    PayrollItem,
    PayrollResult,
    PayrollRun,
)
from app.exporters.excel.f931 import (
    LEGEND,
    SHEET_TITLES,
    build_f931_workbook,
    f931_filename,
)
from app.services.payroll_calculation import calculate_payroll_run
from scripts.seed_demo import seed


@pytest.fixture
def calculated_run(
    tmp_path: Path,
) -> Iterator[tuple[sessionmaker[Session], int, dict[str, Any]]]:
    engine = create_engine(f"sqlite:///{tmp_path / 'f931-export.sqlite'}")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)
    try:
        with factory.begin() as session:
            demo = seed(session)
            admin = demo["admin"]
            company = demo["company"]
            run = PayrollRun(
                organization_id=demo["organization"].id,
                company_id=company.id,
                code="LIQ-2026-06",
                period_start=date(2026, 6, 1),
                period_end=date(2026, 6, 30),
                run_type="MONTHLY",
                status="DRAFT",
                created_by=admin.id,
            )
            session.add(run)
            session.flush()
            calculation = calculate_payroll_run(session, run)
            assert calculation["summary"]["calculated"] == 2
            run_id = run.id

        with factory() as session:
            run = session.get(PayrollRun, run_id)
            assert run is not None
            results = session.scalars(
                select(PayrollResult)
                .where(PayrollResult.payroll_run_id == run_id)
                .order_by(PayrollResult.employee_id)
            ).all()
            items = session.scalars(
                select(PayrollItem)
                .join(PayrollResult, PayrollResult.id == PayrollItem.payroll_result_id)
                .where(PayrollResult.payroll_run_id == run_id)
                .order_by(PayrollResult.employee_id, PayrollItem.id)
            ).all()
            employee_ids = [result.employee_id for result in results]
            employees = session.scalars(
                select(Employee).where(Employee.id.in_(employee_ids)).order_by(Employee.id)
            ).all()
            company = session.get(Company, run.company_id)
            expected = {
                "gross": sum((Decimal(result.gross) for result in results), Decimal("0")),
                "remunerative": sum(
                    (Decimal(result.remunerative) for result in results), Decimal("0")
                ),
                "non_remunerative": sum(
                    (Decimal(result.non_remunerative) for result in results), Decimal("0")
                ),
                "deductions": sum((Decimal(result.deductions) for result in results), Decimal("0")),
                "contributions": sum(
                    (Decimal(result.contributions) for result in results), Decimal("0")
                ),
                "net": sum((Decimal(result.net) for result in results), Decimal("0")),
                "employer_cost": sum(
                    (Decimal(result.employer_cost) for result in results), Decimal("0")
                ),
                "cuils": {employee.cuil for employee in employees},
                "item_count": len(items),
            }
            payload = build_f931_workbook(run, results, items, employees, company)
            expected["payload"] = payload

        yield factory, run_id, expected
    finally:
        engine.dispose()


def _header_columns(worksheet: Any, header_row: int) -> dict[str, int]:
    return {
        str(cell.value): cell.column for cell in worksheet[header_row] if cell.value is not None
    }


def _summary_amounts(worksheet: Any) -> dict[str, float]:
    amounts: dict[str, float] = {}
    for label, amount, _count in worksheet.iter_rows(min_row=1, max_col=3, values_only=True):
        if isinstance(label, str) and isinstance(amount, (int, float)):
            amounts[label] = float(amount)
    return amounts


def test_f931_workbook_has_three_sheets_numeric_totals_and_legend(
    calculated_run: tuple[sessionmaker[Session], int, dict[str, Any]],
) -> None:
    _factory, _run_id, expected = calculated_run
    workbook = load_workbook(io.BytesIO(expected["payload"]), data_only=False)
    try:
        assert workbook.sheetnames == list(SHEET_TITLES)

        for worksheet in workbook.worksheets:
            texts = " ".join(
                str(cell.value)
                for row in worksheet.iter_rows()
                for cell in row
                if cell.value is not None
            )
            assert LEGEND.lower() in texts.lower()

        summary = _summary_amounts(workbook["Resumen"])
        assert summary["Remuneracion bruta"] == pytest.approx(float(expected["gross"]), abs=0.01)
        assert summary["Remuneracion remunerativa"] == pytest.approx(
            float(expected["remunerative"]), abs=0.01
        )
        assert summary["Remuneracion no remunerativa"] == pytest.approx(
            float(expected["non_remunerative"]), abs=0.01
        )
        assert summary["Base imponible total estimada"] == pytest.approx(
            float(expected["remunerative"] + expected["non_remunerative"]), abs=0.01
        )
        assert summary["Aportes y descuentos"] == pytest.approx(
            float(expected["deductions"]), abs=0.01
        )
        assert summary["Contribuciones"] == pytest.approx(
            float(expected["contributions"]), abs=0.01
        )
        assert summary["Neto a pagar"] == pytest.approx(float(expected["net"]), abs=0.01)
        assert summary["Costo total del empleador"] == pytest.approx(
            float(expected["employer_cost"]), abs=0.01
        )
        assert summary["Trabajadores incluidos"] == 2

        workers = workbook["Trabajadores"]
        columns = _header_columns(workers, 4)
        assert {
            "CUIL",
            "Remuneracion bruta",
            "Base imponible total",
            "Aportes",
            "Contribuciones",
        } <= set(columns)
        worker_rows = [
            row
            for row in workers.iter_rows(min_row=5, max_row=workers.max_row, values_only=True)
            if row[0]
        ]
        assert len(worker_rows) == 2
        assert {row[0] for row in worker_rows} == expected["cuils"]
        assert sum(
            float(row[columns["Remuneracion bruta"] - 1]) for row in worker_rows
        ) == pytest.approx(float(expected["gross"]), abs=0.01)
        assert sum(float(row[columns["Aportes"] - 1]) for row in worker_rows) == pytest.approx(
            float(expected["deductions"]), abs=0.01
        )
        for row in worker_rows:
            for column_name in (
                "Remuneracion bruta",
                "Base remunerativa",
                "Base no remunerativa",
                "Base imponible total",
                "Aportes",
                "Contribuciones",
            ):
                assert isinstance(row[columns[column_name] - 1], (int, float))

        concepts = workbook["Conceptos"]
        concept_rows = [
            row
            for row in concepts.iter_rows(min_row=5, max_row=concepts.max_row, values_only=True)
            if row[3]
        ]
        assert len(concept_rows) == expected["item_count"]
        assert all(isinstance(row[8], (int, float)) for row in concept_rows)

        assert all(
            cell.data_type != "f"
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows()
            for cell in row
            if cell.value is not None
        )
    finally:
        workbook.close()


def test_f931_workbook_does_not_execute_user_formulas() -> None:
    run = SimpleNamespace(
        id=7,
        code='..\\LIQ\n"2026',
        period_start=date(2026, 6, 1),
        period_end=date(2026, 6, 30),
        status="CALCULATED",
        run_type="MONTHLY",
    )
    company = SimpleNamespace(legal_name='=HYPERLINK("http://example.com")', cuit="30712345678")
    employee = SimpleNamespace(
        id=1,
        first_name="=1+1",
        last_name='Perez"\n',
        cuil="20123456789",
        employee_code="=BAD",
    )
    result = SimpleNamespace(
        id=10,
        payroll_run_id=7,
        employee_id=1,
        remunerative=Decimal("100.00"),
        non_remunerative=Decimal("0"),
        deductions=Decimal("11.00"),
        contributions=Decimal("20.00"),
        gross=Decimal("100.00"),
        net=Decimal("89.00"),
        employer_cost=Decimal("120.00"),
    )
    item = SimpleNamespace(
        id=1,
        payroll_result_id=10,
        employee_id=1,
        code="=CODE",
        description='=HYPERLINK("http://example.com")',
        nature="REMUNERATIVE",
        quantity=Decimal("1"),
        unit_value=Decimal("100.00"),
        amount=Decimal("100.00"),
    )

    payload = build_f931_workbook(run, [result], [item], [employee], company)
    workbook = load_workbook(io.BytesIO(payload), data_only=False)
    try:
        assert workbook.sheetnames == list(SHEET_TITLES)
        assert workbook["Trabajadores"]["C5"].value == 'Perez", =1+1'
        assert workbook["Conceptos"]["E5"].value == item.description
        assert workbook["Conceptos"]["D5"].value == item.code
        assert all(
            cell.data_type != "f"
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows()
            for cell in row
            if cell.value is not None
        )
    finally:
        workbook.close()

    filename = f931_filename(run)
    assert filename.endswith(".xlsx")
    assert filename.isascii()
    assert "/" not in filename and "\\" not in filename
    assert "\n" not in filename and '"' not in filename

from __future__ import annotations

import io
from collections.abc import Generator
from datetime import date
from pathlib import Path

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from app.api import reports
from app.auth.security import sign_session
from app.config import get_settings
from app.database import Base
from app.database.models import Company, Organization, PayrollRun
from app.database.session import get_db
from app.exporters.excel.f931 import XLSX_MEDIA_TYPE
from app.services.payroll_calculation import calculate_payroll_run
from scripts.seed_demo import seed


def test_f931_endpoint_is_authenticated_scoped_and_requires_calculated_run(
    tmp_path: Path,
) -> None:
    engine = create_engine(f"sqlite:///{tmp_path / 'reports-api.sqlite'}")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)
    try:
        with factory.begin() as session:
            demo = seed(session)
            admin = demo["admin"]
            demo_company = demo["company"]
            calculated = PayrollRun(
                organization_id=demo["organization"].id,
                company_id=demo_company.id,
                code="LIQ-2026-06",
                period_start=date(2026, 6, 1),
                period_end=date(2026, 6, 30),
                run_type="MONTHLY",
                status="DRAFT",
                created_by=admin.id,
            )
            draft = PayrollRun(
                organization_id=demo["organization"].id,
                company_id=demo_company.id,
                code="BORRADOR",
                period_start=date(2026, 7, 1),
                period_end=date(2026, 7, 31),
                run_type="MONTHLY",
                status="DRAFT",
                created_by=admin.id,
            )
            other_organization = Organization(name="Otra Organización", slug="otra")
            session.add_all([calculated, draft, other_organization])
            session.flush()
            other_company = Company(
                organization_id=other_organization.id,
                cuit="30700000001",
                legal_name="Otra Empresa S.A.",
            )
            session.add(other_company)
            session.flush()
            foreign_run = PayrollRun(
                organization_id=other_organization.id,
                company_id=other_company.id,
                code="AJENA",
                period_start=date(2026, 6, 1),
                period_end=date(2026, 6, 30),
                run_type="MONTHLY",
                status="CALCULATED",
                created_by=admin.id,
            )
            session.add(foreign_run)
            session.flush()
            calculation = calculate_payroll_run(session, calculated)
            assert calculation["summary"]["calculated"] == 2
            ids = {
                "admin": admin.id,
                "calculated": calculated.id,
                "draft": draft.id,
                "foreign": foreign_run.id,
            }

        def override_get_db() -> Generator[Session, None, None]:
            with factory() as session:
                yield session

        application = FastAPI()
        application.include_router(reports.router, prefix="/api/v1")
        application.dependency_overrides[get_db] = override_get_db
        try:
            with TestClient(application) as client:
                calculated_url = f"/api/v1/reports/f931/{ids['calculated']}.xlsx"
                assert client.get(calculated_url).status_code == 401

                settings = get_settings()
                client.cookies.set(
                    settings.session_cookie,
                    sign_session(ids["admin"], settings.secret_key),
                )
                response = client.get(calculated_url)
                assert response.status_code == 200, response.text
                assert response.headers["content-type"] == XLSX_MEDIA_TYPE
                disposition = response.headers["content-disposition"]
                assert disposition.startswith('attachment; filename="control-f931-')
                assert disposition.endswith('.xlsx"')
                assert disposition.isascii()

                workbook = load_workbook(io.BytesIO(response.content), data_only=False)
                try:
                    assert workbook.sheetnames == ["Resumen", "Trabajadores", "Conceptos"]
                finally:
                    workbook.close()

                for allowed_status in ("APPROVED", "CLOSED"):
                    with factory.begin() as session:
                        run = session.get(PayrollRun, ids["calculated"])
                        assert run is not None
                        run.status = allowed_status
                    assert client.get(calculated_url).status_code == 200

                foreign = client.get(f"/api/v1/reports/f931/{ids['foreign']}.xlsx")
                assert foreign.status_code == 404

                pending = client.get(f"/api/v1/reports/f931/{ids['draft']}.xlsx")
                assert pending.status_code == 409
                assert "calculada" in pending.json()["detail"].lower()

                missing = client.get("/api/v1/reports/f931/999999.xlsx")
                assert missing.status_code == 404
        finally:
            application.dependency_overrides.pop(get_db, None)
    finally:
        engine.dispose()

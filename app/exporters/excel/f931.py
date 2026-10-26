"""Planilla interna de control F.931 en formato XLSX.

El reporte consolida los totales persistidos de una liquidacion. Es una
herramienta de control interno: no reproduce el diseno oficial del F.931, no
genera la declaracion jurada y no afirma aceptacion por parte de ARCA.

Solo se escriben valores calculados en Python. Ningun texto proveniente de la
base se interpreta como formula de Excel.
"""

from __future__ import annotations

import io
import re
import unicodedata
from collections.abc import Iterable, Mapping
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

LEGEND = "planilla interna de control, no presentacion ARCA"
DISCLAIMER = (
    "Este reporte no reemplaza la declaracion jurada F.931 ni constituye una "
    "presentacion oficial ante ARCA."
)
REPORTABLE_STATUSES = frozenset({"CALCULATED", "APPROVED", "CLOSED"})
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEET_TITLES = ("Resumen", "Trabajadores", "Conceptos")

MONEY_FORMAT = "#,##0.00"
QUANTITY_FORMAT = "#,##0.####"

TITLE_FONT = Font(bold=True, size=14, color="17365D")
LEGEND_FONT = Font(bold=True, size=10, color="9C3D00")
NOTE_FONT = Font(italic=True, size=9, color="5A6577")
SECTION_FONT = Font(bold=True, size=11, color="17365D")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FONT = Font(bold=True, color="17365D")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
LEGEND_FILL = PatternFill("solid", fgColor="FCE4D6")
TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="B7C4D6")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

_NATURE_LABELS = {
    "REMUNERATIVE": "Remunerativo",
    "NON_REMUNERATIVE": "No remunerativo",
    "DEDUCTION": "Descuento",
    "CONTRIBUTION": "Contribucion",
    "AUXILIARY": "Auxiliar",
}
_ILLEGAL_CHARACTERS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _value(source: Any, name: str, default: Any = None) -> Any:
    if source is None:
        return default
    if isinstance(source, Mapping):
        return source.get(name, default)
    return getattr(source, name, default)


def _decimal(value: Any) -> Decimal:
    if value is None or value == "":
        return Decimal("0")
    if isinstance(value, Decimal):
        return value
    try:
        return Decimal(str(value).strip())
    except (InvalidOperation, TypeError, ValueError):
        return Decimal("0")


def _rounded(value: Any) -> Decimal:
    return _decimal(value).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _money(value: Any) -> float:
    return float(_rounded(value))


def _quantity(value: Any) -> float | None:
    if value is None or value == "":
        return None
    return float(_decimal(value).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP))


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return _ILLEGAL_CHARACTERS.sub("", str(value)).strip()


def _slug(value: Any, fallback: str) -> str:
    normalized = unicodedata.normalize("NFKD", _text(value))
    ascii_value = normalized.encode("ascii", "ignore").decode("ascii")
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "-", ascii_value).strip("-._")
    return cleaned[:80] or fallback


def f931_filename(run: Any) -> str:
    """Return an ASCII-only, path-safe filename for the control workbook."""

    code = _slug(_value(run, "code"), "liquidacion")
    period = _value(run, "period_end") or _value(run, "period_start")
    if isinstance(period, (date, datetime)):
        period = period.strftime("%Y%m%d")
    # B-82: marca sin validez ante ARCA en el nombre de la descarga.
    return f"control-f931-{code}-{_slug(period, 'sin-periodo')}-sin-validez-arca.xlsx"


def _period_label(run: Any) -> str:
    start = _value(run, "period_start")
    end = _value(run, "period_end")
    if isinstance(start, (date, datetime)) and isinstance(end, (date, datetime)):
        return f"{start:%d/%m/%Y} al {end:%d/%m/%Y}"
    return _text(end or start)


def _employee_name(employee: Any) -> str:
    first_name = _text(_value(employee, "first_name"))
    last_name = _text(_value(employee, "last_name"))
    if last_name and first_name:
        return f"{last_name}, {first_name}"
    return last_name or first_name


def _employee_index(employees: Any) -> dict[Any, Any]:
    if employees is None:
        return {}
    if isinstance(employees, Mapping):
        index: dict[Any, Any] = {}
        for key, employee in employees.items():
            index[key] = employee
            if isinstance(key, str) and key.isdigit():
                index[int(key)] = employee
            employee_id = _value(employee, "id")
            if employee_id is not None:
                index[employee_id] = employee
        return index
    return {_value(employee, "id"): employee for employee in employees}


def _collect_records(records: Any) -> list[Any]:
    if records is None:
        return []
    if isinstance(records, Mapping):
        return list(records.values())
    return list(records)


def _collect_items(items: Any) -> list[Any]:
    if items is None:
        return []
    if isinstance(items, Mapping):
        collected: list[Any] = []
        for entry in items.values():
            if entry is None:
                continue
            if isinstance(entry, (list, tuple, set)):
                collected.extend(entry)
            else:
                collected.append(entry)
        return collected
    return list(items)


def _write_text(
    worksheet: Worksheet,
    row: int,
    column: int,
    value: Any,
    *,
    font: Font | None = None,
    fill: PatternFill | None = None,
    border: bool = False,
    align: str = "left",
) -> None:
    cell = worksheet.cell(row=row, column=column)
    cell.value = _text(value)
    if cell.data_type == "f":
        # openpyxl marks strings starting with "=" as formulas. Force plain text.
        cell.data_type = "s"
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill
    if border:
        cell.border = CELL_BORDER
    cell.alignment = Alignment(horizontal=align, vertical="center")


def _write_number(
    worksheet: Worksheet,
    row: int,
    column: int,
    value: Any,
    *,
    number_format: str = MONEY_FORMAT,
    font: Font | None = None,
    fill: PatternFill | None = None,
    border: bool = True,
) -> None:
    cell = worksheet.cell(row=row, column=column)
    cell.value = float(value) if isinstance(value, Decimal) else value
    cell.number_format = number_format
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill
    if border:
        cell.border = CELL_BORDER
    cell.alignment = Alignment(horizontal="right", vertical="center")


def _header_row(worksheet: Worksheet, row: int, headers: Iterable[str]) -> None:
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=row, column=column)
        cell.value = header
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = CELL_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[row].height = 30


def _title_block(worksheet: Worksheet, title: str, last_column: int) -> None:
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    _write_text(worksheet, 1, 1, title, font=TITLE_FONT)
    worksheet.row_dimensions[1].height = 24

    worksheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    _write_text(worksheet, 2, 1, LEGEND, font=LEGEND_FONT, fill=LEGEND_FILL)
    worksheet.row_dimensions[2].height = 20


def _fit_columns(worksheet: Worksheet, minimum: int = 11, maximum: int = 44) -> None:
    for cells in worksheet.columns:
        column_index = cells[0].column
        length = max(
            (len(str(cell.value)) for cell in cells if cell.value is not None),
            default=0,
        )
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(length + 2, minimum), maximum
        )


def _build_summary(
    workbook: Workbook,
    run: Any,
    results: list[Any],
    employees: dict[Any, Any],
    company: Any,
    totals: dict[str, Decimal],
) -> None:
    worksheet = workbook.active
    worksheet.title = "Resumen"
    _title_block(worksheet, "Planilla interna de control F.931", 3)
    worksheet.merge_cells("A3:C3")
    _write_text(worksheet, 3, 1, DISCLAIMER, font=NOTE_FONT)

    metadata = (
        ("Empresa", _value(company, "legal_name")),
        ("CUIT", _value(company, "cuit")),
        ("Liquidacion", _value(run, "code")),
        ("Periodo", _period_label(run)),
        ("Estado", _value(run, "status")),
        ("Tipo de liquidacion", _value(run, "run_type")),
    )
    row = 5
    for label, value in metadata:
        _write_text(worksheet, row, 1, label, font=Font(bold=True))
        worksheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        _write_text(worksheet, row, 2, value)
        row += 1

    row += 1
    _write_text(worksheet, row, 1, "Totales de la liquidacion", font=SECTION_FONT)
    row += 1
    _header_row(worksheet, row, ("Concepto", "Importe", "Trabajadores"))
    row += 1

    worker_count = len(results)
    summary_rows = (
        ("Trabajadores incluidos", Decimal(worker_count), worker_count),
        ("Remuneracion bruta", totals["gross"], None),
        ("Remuneracion remunerativa", totals["remunerative"], None),
        ("Remuneracion no remunerativa", totals["non_remunerative"], None),
        ("Base imponible total estimada", totals["base"], None),
        ("Aportes y descuentos", totals["deductions"], None),
        ("Contribuciones", totals["contributions"], None),
        ("Neto a pagar", totals["net"], None),
        ("Costo total del empleador", totals["employer_cost"], None),
    )
    for index, (label, amount, count) in enumerate(summary_rows):
        is_worker_count = label == "Trabajadores incluidos"
        emphasized = index in {0, 1, len(summary_rows) - 1}
        fill = TOTAL_FILL if emphasized else None
        _write_text(
            worksheet,
            row,
            1,
            label,
            font=TOTAL_FONT if emphasized else None,
            fill=fill,
            border=True,
        )
        if is_worker_count:
            _write_number(
                worksheet,
                row,
                2,
                count,
                number_format="0",
                font=TOTAL_FONT if index == 0 else None,
                fill=fill,
            )
            _write_number(
                worksheet,
                row,
                3,
                count,
                number_format="0",
                font=TOTAL_FONT if index == 0 else None,
                fill=fill,
            )
        else:
            _write_number(
                worksheet,
                row,
                2,
                _money(amount),
                font=TOTAL_FONT if emphasized else None,
                fill=fill,
            )
            _write_text(worksheet, row, 3, "", border=True)
        row += 1

    row += 1
    worksheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    _write_text(
        worksheet,
        row,
        1,
        "Bases y obligaciones estimadas a partir de los totales persistidos; "
        "no incluyen apertura fiscal por subsistema ni diferencias contra una "
        "declaracion externa.",
        font=NOTE_FONT,
    )

    worksheet.freeze_panes = "A5"
    worksheet.sheet_view.showGridLines = False
    _fit_columns(worksheet, minimum=18, maximum=46)


def _build_workers(
    workbook: Workbook,
    results: list[Any],
    employees: dict[Any, Any],
) -> None:
    worksheet = workbook.create_sheet("Trabajadores")
    headers = (
        "CUIL",
        "Legajo",
        "Trabajador",
        "Remuneracion bruta",
        "Base remunerativa",
        "Base no remunerativa",
        "Base imponible total",
        "Aportes",
        "Contribuciones",
        "Neto",
        "Costo empleador",
    )
    _title_block(worksheet, "Planilla interna F.931 - Trabajadores", len(headers))
    _header_row(worksheet, 4, headers)

    for row, result in enumerate(results, start=5):
        employee = employees.get(_value(result, "employee_id"))
        remunerative = _rounded(_value(result, "remunerative"))
        non_remunerative = _rounded(_value(result, "non_remunerative"))
        base = remunerative + non_remunerative
        values: tuple[Any, ...] = (
            _text(_value(employee, "cuil")),
            _text(_value(employee, "employee_code")),
            _employee_name(employee),
        )
        for column, value in enumerate(values, start=1):
            _write_text(worksheet, row, column, value, border=True)
        numeric_values = (
            (_money(_value(result, "gross")), MONEY_FORMAT),
            (_money(remunerative), MONEY_FORMAT),
            (_money(non_remunerative), MONEY_FORMAT),
            (_money(base), MONEY_FORMAT),
            (_money(_value(result, "deductions")), MONEY_FORMAT),
            (_money(_value(result, "contributions")), MONEY_FORMAT),
            (_money(_value(result, "net")), MONEY_FORMAT),
            (_money(_value(result, "employer_cost")), MONEY_FORMAT),
        )
        for column, (value, number_format) in enumerate(numeric_values, start=4):
            _write_number(worksheet, row, column, value, number_format=number_format)

    if results:
        worksheet.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{4 + len(results)}"
    worksheet.freeze_panes = "A5"
    worksheet.sheet_view.showGridLines = False
    _fit_columns(worksheet)


def _item_employee_id(item: Any, result_index: dict[Any, Any]) -> Any:
    result = result_index.get(_value(item, "payroll_result_id"))
    return _value(item, "employee_id", _value(result, "employee_id"))


def _build_concepts(
    workbook: Workbook,
    results: list[Any],
    items: list[Any],
    employees: dict[Any, Any],
) -> None:
    worksheet = workbook.create_sheet("Conceptos")
    headers = (
        "CUIL",
        "Legajo",
        "Trabajador",
        "Codigo",
        "Concepto",
        "Naturaleza",
        "Cantidad",
        "Valor unitario",
        "Importe",
    )
    _title_block(worksheet, "Planilla interna F.931 - Conceptos", len(headers))
    _header_row(worksheet, 4, headers)

    result_index = {_value(result, "id"): result for result in results}
    result_ids = set(result_index)
    selected_items: list[Any] = []
    for item in items:
        result_id = _value(item, "payroll_result_id")
        if result_id is None or result_id in result_ids:
            selected_items.append(item)
    selected_items.sort(
        key=lambda item: (
            _text(
                _value(
                    employees.get(_item_employee_id(item, result_index)),
                    "employee_code",
                )
            ),
            _value(item, "id") or 0,
        )
    )

    for row, item in enumerate(selected_items, start=5):
        employee = employees.get(_item_employee_id(item, result_index))
        nature = _text(_value(item, "nature"))
        _write_text(worksheet, row, 1, _value(employee, "cuil"), border=True)
        _write_text(worksheet, row, 2, _value(employee, "employee_code"), border=True)
        _write_text(worksheet, row, 3, _employee_name(employee), border=True)
        _write_text(worksheet, row, 4, _value(item, "code"), border=True)
        _write_text(worksheet, row, 5, _value(item, "description"), border=True)
        _write_text(
            worksheet,
            row,
            6,
            _NATURE_LABELS.get(nature.upper(), nature),
            border=True,
        )
        quantity = _quantity(_value(item, "quantity"))
        if quantity is None:
            _write_text(worksheet, row, 7, "", border=True)
        else:
            _write_number(worksheet, row, 7, quantity, number_format=QUANTITY_FORMAT)
        unit_value = _quantity(_value(item, "unit_value"))
        if unit_value is None:
            _write_text(worksheet, row, 8, "", border=True)
        else:
            _write_number(worksheet, row, 8, unit_value, number_format=MONEY_FORMAT)
        _write_number(worksheet, row, 9, _money(_value(item, "amount")))

    if selected_items:
        worksheet.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{4 + len(selected_items)}"
    worksheet.freeze_panes = "A5"
    worksheet.sheet_view.showGridLines = False
    _fit_columns(worksheet)


def build_f931_workbook(
    run: Any,
    results: Iterable[Any],
    items: Iterable[Any],
    employees: Iterable[Any] | Mapping[Any, Any],
    company: Any,
) -> bytes:
    """Build the internal F.931 control workbook and return its XLSX bytes."""

    run_id = _value(run, "id")
    result_rows = [
        result
        for result in _collect_records(results)
        if run_id is None
        or _value(result, "payroll_run_id") is None
        or _value(result, "payroll_run_id") == run_id
    ]
    employee_index = _employee_index(employees)
    item_rows = _collect_items(items)

    totals = {
        "remunerative": sum(
            (_rounded(_value(result, "remunerative")) for result in result_rows),
            Decimal("0"),
        ),
        "non_remunerative": sum(
            (_rounded(_value(result, "non_remunerative")) for result in result_rows),
            Decimal("0"),
        ),
        "deductions": sum(
            (_rounded(_value(result, "deductions")) for result in result_rows),
            Decimal("0"),
        ),
        "contributions": sum(
            (_rounded(_value(result, "contributions")) for result in result_rows),
            Decimal("0"),
        ),
        "gross": sum(
            (_rounded(_value(result, "gross")) for result in result_rows),
            Decimal("0"),
        ),
        "net": sum(
            (_rounded(_value(result, "net")) for result in result_rows),
            Decimal("0"),
        ),
        "employer_cost": sum(
            (_rounded(_value(result, "employer_cost")) for result in result_rows),
            Decimal("0"),
        ),
    }
    totals["base"] = totals["remunerative"] + totals["non_remunerative"]

    workbook = Workbook()
    workbook.properties.title = "Planilla interna de control F.931"
    workbook.properties.subject = LEGEND
    workbook.properties.creator = "mrbot-payroll"
    workbook.properties.description = DISCLAIMER

    _build_summary(workbook, run, result_rows, employee_index, company, totals)
    _build_workers(workbook, result_rows, employee_index)
    _build_concepts(workbook, result_rows, item_rows, employee_index)

    buffer = io.BytesIO()
    try:
        workbook.save(buffer)
    finally:
        workbook.close()
    return buffer.getvalue()


__all__ = [
    "DISCLAIMER",
    "LEGEND",
    "REPORTABLE_STATUSES",
    "SHEET_TITLES",
    "XLSX_MEDIA_TYPE",
    "build_f931_workbook",
    "f931_filename",
]

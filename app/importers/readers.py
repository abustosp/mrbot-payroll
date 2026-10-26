"""Lectura de archivos CSV y XLSX para importacion.

El resultado es una tabla simple con headers y filas numeradas. No se usan
pandas ni se accede a la red. Los CSV se decodifican probando UTF-8 con BOM,
UTF-8 y Latin-1, y el delimitador se detecta automaticamente.
"""

from __future__ import annotations

import csv
import io
import os
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from app.importers.parsing import is_blank

_CSV_ENCODINGS = ("utf-8-sig", "utf-8", "latin-1")
_DELIMITERS = ",;\t|"


class ImportFileError(ValueError):
    """Error de lectura del archivo de importacion."""


@dataclass
class Table:
    headers: list[str] = field(default_factory=list)
    rows: list[tuple[int, list[Any]]] = field(default_factory=list)
    source: str = "csv"


def _read_source(stream: Any) -> tuple[bytes | str, str]:
    if isinstance(stream, (bytes, bytearray, memoryview)):
        return bytes(stream), ""
    if isinstance(stream, os.PathLike):
        path = Path(stream)
        return path.read_bytes(), path.name
    if isinstance(stream, str):
        # Se acepta una ruta o directamente el contenido CSV.
        try:
            path = Path(stream)
            if path.exists():
                return path.read_bytes(), path.name
        except OSError:
            pass
        return stream, ""
    if hasattr(stream, "read"):
        data = stream.read()
        if isinstance(data, (bytes, bytearray, memoryview)):
            data = bytes(data)
        elif data is None:
            data = ""
        return data, str(getattr(stream, "name", "") or "")
    raise TypeError("Se esperaba bytes, una ruta o un stream con read().")


def detect_delimiter(text: str) -> str:
    """Detecta el delimitador de un CSV a partir de las primeras lineas."""
    sample = "\n".join(text.splitlines()[:5])
    if not sample:
        return ","
    try:
        return csv.Sniffer().sniff(sample, delimiters=_DELIMITERS).delimiter
    except csv.Error:
        counts = {delimiter: sample.count(delimiter) for delimiter in _DELIMITERS}
        best = max(counts, key=lambda delimiter: counts[delimiter])
        return best if counts[best] > 0 else ","


def _decode_csv(data: bytes) -> str:
    for encoding in _CSV_ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ImportFileError("No se pudo decodificar el archivo CSV.")


def _clean_headers(values: list[Any]) -> list[str]:
    headers = ["" if value is None else str(value).strip() for value in values]
    while headers and not headers[-1]:
        headers.pop()
    return headers


def _pad_row(values: list[Any], width: int) -> list[Any]:
    if len(values) < width:
        return values + [None] * (width - len(values))
    return values[:width]


def parse_csv(data: bytes) -> Table:
    text = _decode_csv(data)
    if not text.strip():
        return Table()

    delimiter = detect_delimiter(text)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    headers: list[str] | None = None
    rows: list[tuple[int, list[Any]]] = []

    for values in reader:
        if not values or all(is_blank(value) for value in values):
            continue
        if headers is None:
            headers = _clean_headers(list(values))
            continue
        rows.append((reader.line_num, _pad_row(list(values), len(headers))))

    return Table(headers=headers or [], rows=rows, source="csv")


def _clean_xlsx_headers(values: list[Any]) -> list[str]:
    headers = ["" if value is None else str(value).strip() for value in values]
    while headers and not headers[-1]:
        headers.pop()
    return headers


def parse_xlsx(data: bytes) -> Table:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - la dependencia esta declarada
        raise ImportFileError("openpyxl no esta instalado.") from exc

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl lanza BadZipFile/InvalidFileException
        raise ImportFileError("El archivo XLSX no se pudo leer.") from exc

    headers: list[str] | None = None
    rows: list[tuple[int, list[Any]]] = []
    try:
        worksheet = workbook.active
        for row_number, values in enumerate(worksheet.iter_rows(values_only=True), start=1):
            values = list(values)
            if not values or all(is_blank(value) for value in values):
                continue
            if headers is None:
                headers = _clean_xlsx_headers(values)
                continue
            rows.append((row_number, _pad_row(values, len(headers))))
    finally:
        workbook.close()

    return Table(headers=headers or [], rows=rows, source="xlsx")


def _looks_like_xlsx(data: bytes) -> bool:
    return data.startswith(b"PK\x03\x04")


def read_table(stream: Any) -> Table:
    """Lee un stream CSV/XLSX y devuelve una :class:`Table`."""
    data, name = _read_source(stream)
    suffix = Path(name).suffix.lower()

    if isinstance(data, str):
        if suffix in {".xlsx", ".xlsm"}:
            raise ImportFileError("El XLSX debe abrirse en modo binario.")
        return parse_csv(data.encode("utf-8"))

    if suffix in {".xlsx", ".xlsm"} or _looks_like_xlsx(data):
        return parse_xlsx(data)
    if suffix == ".xls":
        raise ImportFileError("El formato .xls no esta soportado; use .xlsx o CSV.")
    return parse_csv(data)

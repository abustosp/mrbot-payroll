"""Verify pinned ARCA sources and compare synthetic records with its spreadsheets."""
from __future__ import annotations

import argparse
import hashlib
import json
import subprocess
import sys
import tempfile
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse
from urllib.request import urlopen
from zipfile import ZipFile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.exporters.txt.lsd import LAYOUT_VERSION, render_record  # noqa: E402

CATALOG = ROOT / 'normative/lsd/catalog.json'
GOLDEN = ROOT / 'tests/fixtures/lsd_official_golden.json'
SHEETS = {
    '01': ('reg1', 9, 'J9'), '02': ('reg2', 5, 'K5'),
    '03': ('reg3', 5, 'J5'), '04': ('reg4', 5, 'AW5'),
    '05': ('reg5', 5, 'J5'), '06': ('reg 6', 5, 'E5'),
    'concept': ('Armado txt', 8, 'U8'),
}
# Sheet columns C.. (B.. for concepts), independently specified from the exporter.
FIELDS = {
    '01': ('cuit', 'send_id', 'period', 'liquidation_type', 'liquidation_number', 'base_days', 'worker_count'),
    '02': ('cuil', 'employee_id', 'branch', 'bank_account', 'days', 'pay_date', 'rubric_date', 'payment_method'),
    '03': ('cuil', 'concept_code', 'quantity', 'unit', 'amount', 'debit_credit', 'adjustment_period'),
    '04': ('cuil', 'spouse', 'children', 'agreement', 'life_insurance', 'reduction', 'employer_type', 'operation_type', 'employment_status', 'condition', 'activity', 'contract_mode', 'disability', 'locality', 'situation_1', 'start_day_1', 'situation_2', 'start_day_2', 'situation_3', 'start_day_3', 'worked_days', 'worked_hours', 'additional_contribution_rate', 'differential_task_rate', 'health_plan', 'dependents', 'additional_health_contribution', 'additional_health_employer', 'differential_health_base_employee', 'differential_health_base_employer', 'differential_work_risk_base', 'maternity_remuneration', 'gross_remuneration', *(f'taxable_base_{i}' for i in range(1, 10)), 'differential_social_security_employee', 'differential_social_security_employer', 'taxable_base_10', 'deduction_amount'),
    '05': ('cuil', 'category', 'position', 'start_date', 'end_date', 'amount', 'employer_cuit'),
    '06': ('cuil', 'observation'),
    'concept': ('arca_code', 'employer_code', 'description', 'repeat_flag', 'sipa_ap', 'sipa_contrib', 'inssjyp_ap', 'inssjyp_contrib', 'os_ap', 'os_contrib', 'fsr_ap', 'fsr_contrib', 'renatea_ap', 'renatea_contrib', 'aaff_contrib', 'fne_contrib', 'lrt_contrib', 'dif_ap', 'esp_ap'),
}
AMOUNTS_04 = set(FIELDS['04'][22:24]) | set(FIELDS['04'][26:])


def sources(directory: Path | None) -> tuple[dict[str, bytes], list[dict[str, str]]]:
    result, evidence = {}, []
    for document in json.loads(CATALOG.read_text())['documents']:
        url = document['source_url']
        if directory is None:
            with urlopen(url, timeout=60) as response:
                data = response.read()
        else:
            data = (directory / Path(urlparse(url).path).name).read_bytes()
        digest = hashlib.sha256(data).hexdigest()
        if digest != document['sha256']:
            raise ValueError(f"SHA-256 no coincide: {document['id']} ({digest})")
        result[document['id']] = data
        evidence.append({'id': document['id'], 'url': url, 'sha256': digest})
    return result, evidence


def prepare(source_data: dict[str, bytes], cases: list[dict], directory: Path) -> list[Path]:
    catalog = json.loads(CATALOG.read_text())['documents']
    paths = []
    for doc in catalog:
        if doc['doc_type'] != 'herramienta_xlsx':
            continue
        source = directory / doc['inner_file']
        with ZipFile(BytesIO(source_data[doc['id']])) as archive:
            source.write_bytes(archive.read(doc['inner_file']))
        workbook = load_workbook(source)
        relevant = [c for c in cases if (c['record_type'] == 'concept') == (doc['id'] == 'lsd-tool-con')]
        for sheet in workbook:
            if sheet.title in {SHEETS[c['record_type']][0] for c in relevant}:
                # Row 6 in reg3 is retained solely for the adjustment-period probe.
                last = 9 if sheet.title == 'reg1' else 8 if sheet.title == 'Armado txt' else 6 if sheet.title == 'reg3' else 5
                if sheet.max_row > last:
                    sheet.delete_rows(last + 1, sheet.max_row - last)
        for case in relevant:
            kind, values = case['record_type'], case['values']
            sheet_name, row, _ = SHEETS[kind]
            sheet = workbook[sheet_name]
            first_column = 2 if kind == 'concept' else 3
            for column, field in enumerate(FIELDS[kind], start=first_column):
                value = values.get(field, '')
                # TEXT(cell*100) in the official formulas requires numeric cells.
                if ((kind == '04' and field in AMOUNTS_04) or (kind in ('03', '05') and field in ('quantity', 'amount'))) and value != '':
                    value = float(value)
                sheet.cell(row, column).value = value or None
            if kind == '03':
                for column in range(2, 10):
                    sheet.cell(6, column).value = sheet.cell(5, column).value
                sheet['I6'] = '202605'
        workbook.save(source)
        paths.append(source)
    return paths


def calculate(paths: list[Path], directory: Path, executable: str) -> dict[str, object]:
    output = directory / 'output'
    output.mkdir()
    profile = (directory / 'profile').as_uri()
    command = [executable, f'-env:UserInstallation={profile}', '--headless', '--convert-to', 'xlsx', '--outdir', str(output), *(str(p) for p in paths)]
    run = subprocess.run(command, capture_output=True, text=True, timeout=180, check=False)
    if run.returncode or any(not (output / p.name).exists() for p in paths):
        raise RuntimeError(f'LibreOffice no pudo recalcular: {run.stdout} {run.stderr}')
    workbooks = {p.name: load_workbook(output / p.name, data_only=True) for p in paths}
    results = {}
    for kind, (sheet, _, cell) in SHEETS.items():
        name = 'LSD-ARMADO-TXT-Conceptos.xlsx' if kind == 'concept' else 'LSD-ARMADO-TXT-Liquidaciones.xlsx'
        value = workbooks[name][sheet][cell].value
        results[kind] = value
    results['03_adjustment_probe'] = workbooks['LSD-ARMADO-TXT-Liquidaciones.xlsx']['reg3']['J6'].value
    return results


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-dir', type=Path, help='Usar documentos descargados en vez de red')
    parser.add_argument('--recalculate', action='store_true', help='Recalcular planillas con LibreOffice')
    parser.add_argument('--libreoffice', default='libreoffice')
    parser.add_argument('--write-golden', action='store_true', help='Guardar salidas oficiales como fixture, tras comparación exitosa')
    parser.add_argument('--report', type=Path, help='Guardar evidencia JSON')
    args = parser.parse_args()
    if args.write_golden and not args.recalculate:
        parser.error('--write-golden requiere --recalculate')
    source_data, source_evidence = sources(args.source_dir)
    fixture = json.loads(GOLDEN.read_text())
    cases = fixture['cases']
    official = None
    office_version = None
    if args.recalculate:
        office_version = subprocess.run([args.libreoffice, '--version'], capture_output=True, text=True, timeout=30, check=True).stdout.strip()
        with tempfile.TemporaryDirectory(prefix='lsd-official-') as temp:
            paths = prepare(source_data, cases, Path(temp))
            official = calculate(paths, Path(temp), args.libreoffice)
    comparisons = []
    for case in cases:
        kind = case['record_type']
        actual = render_record(kind, case['values'])
        expected = official[kind] if official is not None else case.get('expected')
        comparisons.append({'record_type': kind, 'matched': actual == expected, 'golden_matched': case.get('expected') == expected, 'official_length': len(expected) if isinstance(expected, str) else None, 'exporter_length': len(actual), 'official': expected, 'exporter': actual})
        if official is not None and isinstance(expected, str):
            case['expected'] = expected
    probe = None
    if official is not None:
        value = official['03_adjustment_probe']
        probe = {'input': '202605', 'official_prefix': value[:70] if isinstance(value, str) else value, 'official_length': len(value) if isinstance(value, str) else None, 'expected_length': 51}
    report = {'generated_at_utc': datetime.now(UTC).isoformat(), 'layout': LAYOUT_VERSION, 'sources': source_evidence, 'spreadsheet_recalculated': official is not None, 'expected_source': 'recalculated_spreadsheet' if official is not None else 'golden_fixture', 'libreoffice_version': office_version, 'comparisons': comparisons, 'adjustment_period_probe': probe}
    output = json.dumps(report, ensure_ascii=False, indent=2) + '\n'
    if args.report:
        args.report.write_text(output)
    else:
        print(output, end='')
    success = all(item['matched'] and (args.write_golden or item['golden_matched']) for item in comparisons)
    if args.write_golden and success:
        GOLDEN.write_text(json.dumps(fixture, ensure_ascii=False, indent=2) + '\n')
    return 0 if success else 1


if __name__ == '__main__':
    raise SystemExit(main())
